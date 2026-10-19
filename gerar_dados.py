from __future__ import annotations

import json
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


WORKBOOK_NAME = "Caso AgroInvest - acompanhamento.xlsx"
STAGES = range(1, 6)
CHECKPOINT_SIZE = 5
QUESTION_COLUMNS = {
    "id": "id",
    "stage": "etapa",
    "prompt": "enunciado",
    "alternatives": "alternativas",
    "answer": "gabarito",
    "generalRationale": "por quê",
    "alternativeRationales": "justificativa por alternativa",
}


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).strip().casefold()


def header_index(headers: list[Any], phrase: str) -> int:
    needle = normalize(phrase)
    for index, header in enumerate(headers):
        if needle in normalize(header):
            return index
    raise ValueError(f"Coluna não encontrada: {phrase}")


def parse_labeled_text(value: Any) -> dict[str, str]:
    text = str(value or "").strip()
    pieces = re.split(r"(?m)^\s*([a-z])\)\s*", text, flags=re.IGNORECASE)
    result: dict[str, str] = {}
    for index in range(1, len(pieces), 2):
        letter = pieces[index].strip().lower()
        body = pieces[index + 1].strip()
        result[letter] = body
    return result


def strip_correctness_label(text: str) -> str:
    return re.sub(
        r"^\s*(?:CORRETA|ERRADA)\b\.?\s*",
        "",
        text,
        flags=re.IGNORECASE,
    ).strip()


def question_data(sheet: Any) -> list[dict[str, Any]]:
    headers = [cell.value for cell in sheet[1]]
    indices = {
        "id": header_index(headers, "id"),
        "stage": header_index(headers, "etapa"),
        "prompt": header_index(headers, "enunciado"),
        "alternatives": header_index(headers, "alternativas"),
        "answer": header_index(headers, "gabarito"),
        "generalRationale": header_index(headers, "por quê"),
        "alternativeRationales": header_index(
            headers, "justificativa por alternativa"
        ),
    }
    questions: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[indices["id"]] is None:
            continue
        raw_alternatives = parse_labeled_text(row[indices["alternatives"]])
        raw_rationales = parse_labeled_text(row[indices["alternativeRationales"]])
        alternatives = [
            {"letter": letter, "text": text}
            for letter, text in raw_alternatives.items()
        ]
        rationales = {
            letter: strip_correctness_label(text)
            for letter, text in raw_rationales.items()
        }
        question = {
            "id": int(row[indices["id"]]),
            "stage": int(row[indices["stage"]]),
            "prompt": str(row[indices["prompt"]] or "").strip(),
            "alternatives": alternatives,
            "answer": str(row[indices["answer"]] or "").strip().lower(),
            "generalRationale": str(
                row[indices["generalRationale"]] or ""
            ).strip(),
            "rationales": rationales,
        }
        if not question["prompt"] or len(alternatives) < 2:
            raise ValueError(
                f"Questão {question['id']} sem enunciado ou alternativas."
            )
        if question["answer"] not in {item["letter"] for item in alternatives}:
            raise ValueError(
                f"Gabarito inválido na questão {question['id']}."
            )
        questions.append(question)

    for stage in STAGES:
        count = sum(question["stage"] == stage for question in questions)
        if count != CHECKPOINT_SIZE:
            raise ValueError(
                f"A etapa {stage} precisa ter exatamente {CHECKPOINT_SIZE} "
                f"questões; a planilha contém {count}."
            )
    return questions


def company_data(sheet: Any) -> list[dict[str, Any]]:
    headers = [cell.value for cell in sheet[1]]
    records = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        records.append(
            {
                "id": int(row[0]),
                "name": row[1],
                "sector": row[2],
                "biologicalAsset": row[3],
                "bearerPlantAndPpe": row[4],
                "agriculturalProduct": row[5],
                "investmentPropertyStatus": row[6],
                "propertyEvent": row[7],
                "cpcItems": row[8],
                "classReferences": row[9],
                "reportEffect": row[10],
                "situationCount": row[11],
                "group": row[12],
                "status": row[13] if len(row) > 13 else None,
            }
        )
    if not records:
        raise ValueError("A aba Empresas não contém empresas.")
    return records


def find_company_id(name: str, companies: list[dict[str, Any]]) -> int:
    candidate = normalize(name)
    matches = [
        company
        for company in companies
        if normalize(company["name"]).startswith(candidate)
        or candidate.startswith(normalize(company["name"]))
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Não foi possível associar o gabarito à empresa: {name!r}."
        )
    return matches[0]["id"]


def balance_data(
    workbook: Any, companies: list[dict[str, Any]]
) -> tuple[dict[str, Any], list[str]]:
    balances: dict[str, Any] = {}
    missing_sheets = []
    for stage in range(2, 5):
        sheet_name = f"Gabarito E{stage}"
        if sheet_name not in workbook.sheetnames:
            missing_sheets.append(sheet_name)
            continue

        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        if len(headers) < 2:
            raise ValueError(f"Aba {sheet_name} sem cabeçalhos válidos.")
        by_company: dict[str, Any] = {}
        current_company: str | None = None
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0]:
                candidate = str(row[0]).strip()
                try:
                    find_company_id(candidate, companies)
                except ValueError:
                    current_company = None
                    continue
                current_company = candidate
            if not current_company:
                continue
            values = [str(value) for value in row if value is not None]
            if not any(
                re.search(
                    r"saldo|caixa final|resultado acumulado",
                    value,
                    re.IGNORECASE,
                )
                for value in values
            ):
                continue
            company_id = str(find_company_id(current_company, companies))
            fields = {
                str(header): value
                for header, value in zip(headers, row)
                if header is not None and value is not None
            }
            by_company.setdefault(
                company_id,
                {
                    "companyNameInSheet": current_company,
                    "summaryRows": [],
                },
            )["summaryRows"].append(fields)
            if stage == 2 and normalize(row[2]) == "caixa final":
                by_company[company_id]["checkpoint"] = {
                    "cash": row[4],
                    "result": row[5],
                }
            elif stage in (3, 4) and normalize(row[2]).startswith("caixa final"):
                by_company[company_id]["checkpoint"] = {
                    "result": row[5],
                    "cash": row[6],
                }
        balances[str(stage)] = {
            "sourceSheet": sheet_name,
            "companies": by_company,
        }
    return balances, missing_sheets


def stage_one_data(
    workbook: Any, companies: list[dict[str, Any]]
) -> tuple[dict[str, Any], list[str]]:
    sheet_name = "Gabarito E1"
    if sheet_name not in workbook.sheetnames:
        return {}, [sheet_name]

    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    indices = {
        "company": header_index(headers, "empresa"),
        "item": header_index(headers, "item"),
        "asset": header_index(headers, "bem"),
        "classification": header_index(headers, "classificação"),
        "balanceGroup": header_index(headers, "grupo do balanço"),
        "standard": header_index(headers, "norma e item"),
    }
    by_company: dict[str, list[dict[str, Any]]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name = row[indices["company"]]
        if not name or not row[indices["item"]]:
            continue
        company_id = str(find_company_id(str(name).strip(), companies))
        by_company.setdefault(company_id, []).append({
            "item": str(row[indices["item"]]).strip(),
            "asset": str(row[indices["asset"]] or "").strip(),
            "classification": str(row[indices["classification"]] or "").strip(),
            "balanceGroup": str(row[indices["balanceGroup"]] or "").strip(),
            "standard": str(row[indices["standard"]] or "").strip(),
        })
    return by_company, []


def stage_five_data(
    workbook: Any, companies: list[dict[str, Any]]
) -> tuple[dict[str, Any], list[str]]:
    sheet_name = "Gabarito E5"
    if sheet_name not in workbook.sheetnames:
        return {}, [sheet_name]

    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    company_columns: dict[str, int] = {}
    for column_index, header in enumerate(headers[1:], start=1):
        if not header:
            continue
        match = re.match(r"\s*(\d+)\s*[·.]", str(header))
        if not match:
            continue
        company_id = int(match.group(1))
        if not any(company["id"] == company_id for company in companies):
            raise ValueError(
                f"Empresa inválida no cabeçalho de {sheet_name}: {header!r}."
            )
        company_columns[str(company_id)] = column_index

    rows_by_label: dict[str, dict[str, Any]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        label = row[0]
        if not label:
            continue
        values = {
            company_id: row[column_index]
            for company_id, column_index in company_columns.items()
            if column_index < len(row) and row[column_index] is not None
        }
        rows_by_label[str(label).strip()] = values

    dre_rows = []
    for row_number in range(2, sheet.max_row + 1):
        label = sheet.cell(row_number, 1).value
        if not label:
            continue
        if normalize(label).startswith("laudo"):
            break
        if normalize(label).startswith("4."):
            continue
        values = rows_by_label.get(str(label).strip(), {})
        if values:
            dre_rows.append({"label": str(label).strip(), "values": values})

    metric_labels = {
        "laudoA": "1. Efeito no resultado com A",
        "laudoB": "2. Efeito no resultado com B",
        "cashEffect": "3. Entrada de caixa (A e B)",
    }
    checkpoint: dict[str, Any] = {}
    for company in companies:
        company_id = str(company["id"])
        values = {
            key: rows_by_label.get(label, {}).get(company_id)
            for key, label in metric_labels.items()
        }
        if all(value is not None for value in values.values()):
            checkpoint[company_id] = values

    return {
        "sourceSheet": sheet_name,
        "dreRows": dre_rows,
        "checkpoint": checkpoint,
    }, []


def report_blocks_data(
    workbook: Any, companies: list[dict[str, Any]]
) -> tuple[dict[str, list[dict[str, Any]]], list[str]]:
    sheet_name = "Parecer (blocos)"
    if sheet_name not in workbook.sheetnames:
        return {}, [sheet_name]

    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    indices = {
        "company": header_index(headers, "empresa"),
        "block": header_index(headers, "bloco"),
        "theme": header_index(headers, "tema"),
        "option": header_index(headers, "opção"),
        "text": header_index(headers, "texto"),
        "correct": header_index(headers, "correta"),
        "rationale": header_index(headers, "por quê"),
    }
    by_company: dict[str, list[dict[str, Any]]] = {}
    current_company: str | None = None
    current_block: int | None = None
    current_theme: str | None = None
    grouped: dict[str, dict[int, dict[str, Any]]] = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[indices["company"]]:
            current_company = str(row[indices["company"]]).strip()
            company_id = str(find_company_id(current_company, companies))
            grouped.setdefault(company_id, {})
            current_block = None
            current_theme = None
        if row[indices["block"]] is not None:
            current_block = int(row[indices["block"]])
        if row[indices["theme"]]:
            current_theme = str(row[indices["theme"]]).strip()
        if not current_company or current_block is None:
            continue

        company_id = str(find_company_id(current_company, companies))
        block = grouped[company_id].setdefault(
            current_block,
            {
                "id": current_block,
                "theme": current_theme or "",
                "options": [],
            },
        )
        if current_theme and not block["theme"]:
            block["theme"] = current_theme

        text = str(row[indices["text"]] or "").strip()
        option = row[indices["option"]]
        if current_block == 5:
            if text and not block.get("prompt"):
                block["prompt"] = text
            continue
        if option is None or not text:
            continue
        correctness = str(row[indices["correct"]] or "").strip()
        block["options"].append({
            "id": str(option).strip().lower(),
            "text": text,
            "correctness": correctness,
            "correct": normalize(correctness) == "correta",
            "rationale": str(row[indices["rationale"]] or "").strip(),
        })

    for company_id, blocks in grouped.items():
        ordered = [blocks[key] for key in sorted(blocks)]
        for block in ordered:
            if block["id"] in (1, 2, 3):
                correct_count = sum(option["correct"] for option in block["options"])
                if correct_count != 1:
                    raise ValueError(
                        f"O bloco técnico {block['id']} da empresa {company_id} "
                        f"precisa ter exatamente uma opção correta."
                    )
            if block["id"] == 4 and len(block["options"]) < 2:
                raise ValueError(
                    f"O bloco 4 da empresa {company_id} precisa ter duas opções defensáveis."
                )
        by_company[company_id] = ordered
    return by_company, []


def main() -> None:
    output_path = Path(__file__).resolve().parent / "dados.json"
    workbook_candidates = [
        output_path.parent.parent / WORKBOOK_NAME,
        output_path.parent / WORKBOOK_NAME,
    ]
    workbook_path = next(
        (candidate for candidate in workbook_candidates if candidate.is_file()),
        None,
    )
    if workbook_path is None:
        searched = "\n".join(f"- {path}" for path in workbook_candidates)
        raise FileNotFoundError(
            f"Não encontrei {WORKBOOK_NAME}. Caminhos verificados:\n{searched}"
        )

    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    required_sheets = ["Questões", "Empresas"]
    missing_required = [name for name in required_sheets if name not in workbook.sheetnames]
    if missing_required:
        raise ValueError(
            "Abas obrigatórias ausentes: " + ", ".join(missing_required)
        )

    companies = company_data(workbook["Empresas"])
    questions = question_data(workbook["Questões"])
    balances, missing_balance_sheets = balance_data(workbook, companies)
    stage_one, missing_stage_one = stage_one_data(workbook, companies)
    stage_five, missing_stage_five = stage_five_data(workbook, companies)
    report_blocks, missing_report_blocks = report_blocks_data(workbook, companies)
    missing_data_sheets = (
        missing_balance_sheets + missing_stage_one + missing_stage_five
    )
    data = {
        "sourceWorkbook": workbook_path.name,
        "generatedAt": datetime.now().astimezone().isoformat(),
        "companies": companies,
        "questions": questions,
        "balances": balances,
        "stageOneItems": stage_one,
        "stageFive": stage_five,
        "reportBlocks": report_blocks,
        "missingBalanceSheets": missing_data_sheets,
        "missingReportSheets": missing_report_blocks,
    }
    output_path.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(
        f"Gerado {output_path.name}: {len(companies)} empresas, "
        f"{len(questions)} questões e saldos nas etapas "
        f"{', '.join(balances) or 'nenhuma'}."
    )
    if missing_data_sheets:
        print(
            "Abas de dados ausentes: "
            + ", ".join(missing_data_sheets),
            file=sys.stderr,
        )
    if missing_report_blocks:
        print("Abas de parecer ausentes: " + ", ".join(missing_report_blocks), file=sys.stderr)


if __name__ == "__main__":
    main()
